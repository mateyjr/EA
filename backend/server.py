from dotenv import load_dotenv
from pathlib import Path
ROOT_DIR = Path(__file__).parent
load_dotenv(ROOT_DIR / '.env')

import os
import uuid
import logging
import bcrypt
import jwt
from datetime import datetime, timezone, timedelta
from typing import List, Optional, Dict, Any, Literal
from fastapi import FastAPI, APIRouter, HTTPException, Depends, Request, Response, UploadFile, File, Form
from fastapi.security import HTTPBearer
from fastapi.responses import StreamingResponse
from starlette.middleware.cors import CORSMiddleware
import io
import csv
import base64
import httpx
import requests
import hmac
import secrets
from motor.motor_asyncio import AsyncIOMotorClient
from pydantic import BaseModel, Field, EmailStr, ConfigDict

# ---------- Setup ----------
mongo_url = os.environ['MONGO_URL']
client = AsyncIOMotorClient(mongo_url)
db = client[os.environ['DB_NAME']]

JWT_SECRET = os.environ.get('JWT_SECRET', 'change-me')
JWT_ALGO = "HS256"

app = FastAPI(title="Colecle EAMS API")
api = APIRouter(prefix="/api")

logging.basicConfig(level=logging.INFO)
logger = logging.getLogger("eams")

# ---------- Auth helpers ----------
def hash_pw(p: str) -> str:
    return bcrypt.hashpw(p.encode(), bcrypt.gensalt()).decode()

def verify_pw(p: str, h: str) -> bool:
    try:
        return bcrypt.checkpw(p.encode(), h.encode())
    except Exception:
        return False

def make_token(uid: str, email: str, role: str, minutes: int = 60 * 24 * 7) -> str:
    return jwt.encode({
        "sub": uid, "email": email, "role": role,
        "exp": datetime.now(timezone.utc) + timedelta(minutes=minutes),
        "type": "access",
    }, JWT_SECRET, algorithm=JWT_ALGO)

async def get_current_user(request: Request) -> dict:
    token = request.cookies.get("access_token")
    if not token:
        h = request.headers.get("Authorization", "")
        if h.startswith("Bearer "):
            token = h[7:]
    if not token:
        raise HTTPException(status_code=401, detail="Not authenticated")
    try:
        payload = jwt.decode(token, JWT_SECRET, algorithms=[JWT_ALGO])
    except jwt.ExpiredSignatureError:
        raise HTTPException(status_code=401, detail="Token expired")
    except jwt.InvalidTokenError:
        raise HTTPException(status_code=401, detail="Invalid token")
    user = await db.users.find_one({"id": payload["sub"]}, {"_id": 0, "password_hash": 0})
    if not user:
        raise HTTPException(status_code=401, detail="User not found")
    return user

def require_role(*roles: str):
    async def dep(user: dict = Depends(get_current_user)):
        if user.get("role") not in roles and "admin" not in [user.get("role")]:
            raise HTTPException(status_code=403, detail="Insufficient permissions")
        return user
    return dep

# ---------- Models ----------
class UserPublic(BaseModel):
    id: str
    email: str
    name: str
    role: str

class RegisterIn(BaseModel):
    email: EmailStr
    password: str
    name: str
    role: str = "viewer"

class LoginIn(BaseModel):
    email: EmailStr
    password: str

DOMAINS = ["business", "application", "data", "security", "integration", "technology"]

class ArchObject(BaseModel):
    model_config = ConfigDict(extra="ignore")
    id: str = Field(default_factory=lambda: str(uuid.uuid4()))
    domain: str
    type: str  # e.g. capability, process, application, data_entity, integration, control, technology
    code: Optional[str] = None
    name: str
    description: Optional[str] = ""
    owner: Optional[str] = ""
    status: Optional[str] = "Active"  # Draft/Active/Deprecated/Retired
    lifecycle: Optional[str] = "Production"
    criticality: Optional[str] = "Medium"  # Low/Medium/High/Mission-Critical
    tags: List[str] = []
    attributes: Dict[str, Any] = {}
    created_at: str = Field(default_factory=lambda: datetime.now(timezone.utc).isoformat())
    updated_at: str = Field(default_factory=lambda: datetime.now(timezone.utc).isoformat())
    created_by: Optional[str] = None

class ObjectIn(BaseModel):
    domain: str
    type: str
    code: Optional[str] = None
    name: str
    description: Optional[str] = ""
    owner: Optional[str] = ""
    status: Optional[str] = "Active"
    lifecycle: Optional[str] = "Production"
    criticality: Optional[str] = "Medium"
    tags: List[str] = []
    attributes: Dict[str, Any] = {}

class Relationship(BaseModel):
    id: str = Field(default_factory=lambda: str(uuid.uuid4()))
    source_id: str
    target_id: str
    rel_type: str = "supports"  # supports/uses/hosts/consumes/produces/protects/depends_on
    description: Optional[str] = ""
    criticality: Optional[str] = "Medium"
    created_at: str = Field(default_factory=lambda: datetime.now(timezone.utc).isoformat())

class RelationshipIn(BaseModel):
    source_id: str
    target_id: str
    rel_type: str = "supports"
    description: Optional[str] = ""
    criticality: Optional[str] = "Medium"

class ADR(BaseModel):
    id: str = Field(default_factory=lambda: str(uuid.uuid4()))
    number: int
    title: str
    domain: str
    context: str = ""
    problem: str = ""
    options: str = ""
    decision: str = ""
    rationale: str = ""
    consequences: str = ""
    status: str = "Proposed"  # Proposed/Accepted/Superseded/Rejected/Deprecated
    owner: str = ""
    related_objects: List[str] = []
    created_at: str = Field(default_factory=lambda: datetime.now(timezone.utc).isoformat())

class ADRIn(BaseModel):
    title: str
    domain: str
    context: str = ""
    problem: str = ""
    options: str = ""
    decision: str = ""
    rationale: str = ""
    consequences: str = ""
    status: str = "Proposed"
    owner: str = ""
    related_objects: List[str] = []

class Standard(BaseModel):
    id: str = Field(default_factory=lambda: str(uuid.uuid4()))
    code: str
    name: str
    domain: str
    description: str = ""
    mandatory: bool = True
    version: str = "1.0"
    owner: str = ""
    status: str = "Active"
    created_at: str = Field(default_factory=lambda: datetime.now(timezone.utc).isoformat())

class StandardIn(BaseModel):
    code: str
    name: str
    domain: str
    description: str = ""
    mandatory: bool = True
    version: str = "1.0"
    owner: str = ""
    status: str = "Active"

class Risk(BaseModel):
    id: str = Field(default_factory=lambda: str(uuid.uuid4()))
    code: str
    title: str
    domain: str
    related_object: Optional[str] = None
    description: str = ""
    likelihood: str = "Medium"
    impact: str = "Medium"
    score: int = 6
    severity: str = "Medium"
    mitigation: str = ""
    owner: str = ""
    target_date: Optional[str] = None
    status: str = "Open"
    created_at: str = Field(default_factory=lambda: datetime.now(timezone.utc).isoformat())

class RiskIn(BaseModel):
    code: str
    title: str
    domain: str
    related_object: Optional[str] = None
    description: str = ""
    likelihood: str = "Medium"
    impact: str = "Medium"
    score: int = 6
    severity: str = "Medium"
    mitigation: str = ""
    owner: str = ""
    target_date: Optional[str] = None
    status: str = "Open"

class Review(BaseModel):
    id: str = Field(default_factory=lambda: str(uuid.uuid4()))
    code: str
    title: str
    submitter: str
    summary: str = ""
    domains_impacted: List[str] = []
    related_objects: List[str] = []
    stage: str = "Draft"  # Draft/Submitted/Domain Reviews/EA Review/Committee/Approved/Rejected
    findings: List[Dict[str, Any]] = []
    created_at: str = Field(default_factory=lambda: datetime.now(timezone.utc).isoformat())

class ReviewIn(BaseModel):
    title: str
    submitter: str
    summary: str = ""
    domains_impacted: List[str] = []
    related_objects: List[str] = []
    stage: str = "Draft"

class BrandSettings(BaseModel):
    name: str = "Colecle"
    subtitle: str = "System EAMS"
    logo_url: Optional[str] = None
    accent: str = "#f59e0b"
    updated_at: str = Field(default_factory=lambda: datetime.now(timezone.utc).isoformat())

# ---------- Auth Endpoints ----------
@api.post("/auth/register")
async def register(body: RegisterIn, response: Response):
    email = body.email.lower()
    if await db.users.find_one({"email": email}):
        raise HTTPException(status_code=400, detail="Email already registered")
    uid = str(uuid.uuid4())
    doc = {
        "id": uid, "email": email, "password_hash": hash_pw(body.password),
        "name": body.name, "role": body.role or "viewer",
        "created_at": datetime.now(timezone.utc).isoformat(),
    }
    await db.users.insert_one(doc)
    token = make_token(uid, email, doc["role"])
    response.set_cookie("access_token", token, httponly=True, secure=True, samesite="none", max_age=60*60*24*7, path="/")
    return {"id": uid, "email": email, "name": body.name, "role": doc["role"], "token": token}

@api.post("/auth/login")
async def login(body: LoginIn, response: Response):
    email = body.email.lower()
    user = await db.users.find_one({"email": email})
    if not user or not verify_pw(body.password, user["password_hash"]):
        raise HTTPException(status_code=401, detail="Invalid email or password")
    token = make_token(user["id"], email, user["role"])
    response.set_cookie("access_token", token, httponly=True, secure=True, samesite="none", max_age=60*60*24*7, path="/")
    return {"id": user["id"], "email": email, "name": user["name"], "role": user["role"], "token": token}

@api.post("/auth/logout")
async def logout(response: Response):
    response.delete_cookie("access_token", path="/")
    return {"ok": True}

@api.get("/auth/me")
async def me(user: dict = Depends(get_current_user)):
    return {"id": user["id"], "email": user["email"], "name": user["name"], "role": user["role"]}

# ---------- Brand ----------
@api.get("/brand")
async def get_brand():
    b = await db.brand.find_one({"_id": "singleton"}, {"_id": 0})
    if not b:
        b = BrandSettings().model_dump()
        await db.brand.insert_one({"_id": "singleton", **b})
    return b

@api.put("/brand")
async def update_brand(body: BrandSettings, user: dict = Depends(get_current_user)):
    if user["role"] not in ("admin", "lead_architect"):
        raise HTTPException(status_code=403, detail="Admin only")
    body.updated_at = datetime.now(timezone.utc).isoformat()
    await db.brand.update_one({"_id": "singleton"}, {"$set": body.model_dump()}, upsert=True)
    return body

# ---------- Objects (unified catalog) ----------
@api.get("/objects")
async def list_objects(domain: Optional[str] = None, type: Optional[str] = None, q: Optional[str] = None, user: dict = Depends(get_current_user)):
    query: Dict[str, Any] = {}
    if domain: query["domain"] = domain
    if type: query["type"] = type
    if q:
        query["$or"] = [{"name": {"$regex": q, "$options": "i"}}, {"code": {"$regex": q, "$options": "i"}}, {"description": {"$regex": q, "$options": "i"}}]
    items = await db.objects.find(query, {"_id": 0}).sort("name", 1).to_list(2000)
    return items

@api.get("/objects/{oid}")
async def get_object(oid: str, user: dict = Depends(get_current_user)):
    obj = await db.objects.find_one({"id": oid}, {"_id": 0})
    if not obj:
        raise HTTPException(status_code=404, detail="Not found")
    # Upstream (things that link INTO this)
    up_rels = await db.relationships.find({"target_id": oid}, {"_id": 0}).to_list(500)
    down_rels = await db.relationships.find({"source_id": oid}, {"_id": 0}).to_list(500)
    up_ids = [r["source_id"] for r in up_rels]
    down_ids = [r["target_id"] for r in down_rels]
    up_objs = await db.objects.find({"id": {"$in": up_ids}}, {"_id": 0}).to_list(500)
    down_objs = await db.objects.find({"id": {"$in": down_ids}}, {"_id": 0}).to_list(500)
    return {"object": obj, "upstream": up_objs, "downstream": down_objs, "upstream_rels": up_rels, "downstream_rels": down_rels}

@api.post("/objects")
async def create_object(body: ObjectIn, user: dict = Depends(get_current_user)):
    if body.domain not in DOMAINS:
        raise HTTPException(status_code=400, detail="Invalid domain")
    obj = ArchObject(**body.model_dump(), created_by=user["id"])
    await db.objects.insert_one(obj.model_dump())
    await audit(user, "create", obj.id, body.domain, None, obj.model_dump())
    return obj.model_dump()

@api.put("/objects/{oid}")
async def update_object(oid: str, body: ObjectIn, user: dict = Depends(get_current_user)):
    existing = await db.objects.find_one({"id": oid}, {"_id": 0})
    if not existing:
        raise HTTPException(status_code=404, detail="Not found")
    patch = body.model_dump()
    patch["updated_at"] = datetime.now(timezone.utc).isoformat()
    await db.objects.update_one({"id": oid}, {"$set": patch})
    updated = await db.objects.find_one({"id": oid}, {"_id": 0})
    await audit(user, "update", oid, body.domain, existing, updated)
    return updated

@api.delete("/objects/{oid}")
async def delete_object(oid: str, user: dict = Depends(get_current_user)):
    existing = await db.objects.find_one({"id": oid}, {"_id": 0})
    if not existing:
        raise HTTPException(status_code=404, detail="Not found")
    await db.objects.delete_one({"id": oid})
    await db.relationships.delete_many({"$or": [{"source_id": oid}, {"target_id": oid}]})
    await audit(user, "delete", oid, existing.get("domain", ""), existing, None)
    return {"ok": True}

# ---------- Relationships ----------
@api.get("/relationships")
async def list_relationships(user: dict = Depends(get_current_user)):
    return await db.relationships.find({}, {"_id": 0}).to_list(5000)

@api.post("/relationships")
async def create_relationship(body: RelationshipIn, user: dict = Depends(get_current_user)):
    rel = Relationship(**body.model_dump())
    await db.relationships.insert_one(rel.model_dump())
    return rel.model_dump()

@api.delete("/relationships/{rid}")
async def delete_relationship(rid: str, user: dict = Depends(get_current_user)):
    await db.relationships.delete_one({"id": rid})
    return {"ok": True}

# ---------- Traceability & Impact ----------
@api.get("/traceability/{oid}")
async def traceability(oid: str, direction: Literal["forward", "reverse", "both"] = "both", depth: int = 4, user: dict = Depends(get_current_user)):
    root = await db.objects.find_one({"id": oid}, {"_id": 0})
    if not root:
        raise HTTPException(status_code=404, detail="Not found")
    all_rels = await db.relationships.find({}, {"_id": 0}).to_list(10000)
    all_objs = await db.objects.find({}, {"_id": 0}).to_list(10000)
    by_id = {o["id"]: o for o in all_objs}

    def walk(start: str, dirn: str):
        visited = {start}
        edges = []
        frontier = [(start, 0)]
        while frontier:
            cur, d = frontier.pop(0)
            if d >= depth: continue
            for r in all_rels:
                nxt = None
                if dirn == "forward" and r["source_id"] == cur:
                    nxt = r["target_id"]
                elif dirn == "reverse" and r["target_id"] == cur:
                    nxt = r["source_id"]
                if nxt:
                    edges.append(r)
                    if nxt not in visited:
                        visited.add(nxt)
                        frontier.append((nxt, d + 1))
        nodes = [by_id[i] for i in visited if i in by_id]
        return nodes, edges

    nodes_set = {oid: root}
    edges_all = []
    if direction in ("forward", "both"):
        n, e = walk(oid, "forward")
        for nn in n: nodes_set[nn["id"]] = nn
        edges_all.extend(e)
    if direction in ("reverse", "both"):
        n, e = walk(oid, "reverse")
        for nn in n: nodes_set[nn["id"]] = nn
        edges_all.extend(e)
    # de-dup edges
    seen = set()
    edges_dedup = []
    for r in edges_all:
        if r["id"] not in seen:
            seen.add(r["id"])
            edges_dedup.append(r)
    return {"root": root, "nodes": list(nodes_set.values()), "edges": edges_dedup}

@api.get("/impact/{oid}")
async def impact(oid: str, user: dict = Depends(get_current_user)):
    """Reverse impact: what gets broken if this object fails."""
    trace = await traceability(oid, direction="reverse", depth=8, user=user)
    by_domain: Dict[str, List[Dict[str, Any]]] = {d: [] for d in DOMAINS}
    for n in trace["nodes"]:
        if n["id"] == oid: continue
        by_domain.setdefault(n["domain"], []).append(n)
    return {"root": trace["root"], "affected_by_domain": by_domain, "total_affected": sum(len(v) for v in by_domain.values())}

# ---------- Search ----------
@api.get("/search")
async def global_search(q: str, user: dict = Depends(get_current_user)):
    if not q or len(q) < 2:
        return {"results": []}
    regex = {"$regex": q, "$options": "i"}
    objs = await db.objects.find({"$or": [{"name": regex}, {"code": regex}, {"description": regex}]}, {"_id": 0}).limit(30).to_list(30)
    adrs = await db.adrs.find({"$or": [{"title": regex}, {"decision": regex}]}, {"_id": 0}).limit(10).to_list(10)
    stds = await db.standards.find({"$or": [{"name": regex}, {"code": regex}]}, {"_id": 0}).limit(10).to_list(10)
    risks = await db.risks.find({"$or": [{"title": regex}, {"code": regex}]}, {"_id": 0}).limit(10).to_list(10)
    return {"objects": objs, "adrs": adrs, "standards": stds, "risks": risks}

# ---------- ADRs ----------
@api.get("/adrs")
async def list_adrs(user: dict = Depends(get_current_user)):
    return await db.adrs.find({}, {"_id": 0}).sort("number", -1).to_list(500)

@api.post("/adrs")
async def create_adr(body: ADRIn, user: dict = Depends(get_current_user)):
    count = await db.adrs.count_documents({})
    adr = ADR(number=count + 1, **body.model_dump())
    await db.adrs.insert_one(adr.model_dump())
    return adr.model_dump()

@api.put("/adrs/{aid}")
async def update_adr(aid: str, body: ADRIn, user: dict = Depends(get_current_user)):
    await db.adrs.update_one({"id": aid}, {"$set": body.model_dump()})
    return await db.adrs.find_one({"id": aid}, {"_id": 0})

@api.delete("/adrs/{aid}")
async def delete_adr(aid: str, user: dict = Depends(get_current_user)):
    await db.adrs.delete_one({"id": aid})
    return {"ok": True}

# ---------- Standards ----------
@api.get("/standards")
async def list_standards(user: dict = Depends(get_current_user)):
    return await db.standards.find({}, {"_id": 0}).sort("code", 1).to_list(500)

@api.post("/standards")
async def create_standard(body: StandardIn, user: dict = Depends(get_current_user)):
    s = Standard(**body.model_dump())
    await db.standards.insert_one(s.model_dump())
    return s.model_dump()

@api.put("/standards/{sid}")
async def update_standard(sid: str, body: StandardIn, user: dict = Depends(get_current_user)):
    await db.standards.update_one({"id": sid}, {"$set": body.model_dump()})
    return await db.standards.find_one({"id": sid}, {"_id": 0})

@api.delete("/standards/{sid}")
async def delete_standard(sid: str, user: dict = Depends(get_current_user)):
    await db.standards.delete_one({"id": sid})
    return {"ok": True}

# ---------- Risks ----------
@api.get("/risks")
async def list_risks(user: dict = Depends(get_current_user)):
    return await db.risks.find({}, {"_id": 0}).sort("code", 1).to_list(500)

@api.post("/risks")
async def create_risk(body: RiskIn, user: dict = Depends(get_current_user)):
    r = Risk(**body.model_dump())
    await db.risks.insert_one(r.model_dump())
    return r.model_dump()

@api.put("/risks/{rid}")
async def update_risk(rid: str, body: RiskIn, user: dict = Depends(get_current_user)):
    await db.risks.update_one({"id": rid}, {"$set": body.model_dump()})
    return await db.risks.find_one({"id": rid}, {"_id": 0})

@api.delete("/risks/{rid}")
async def delete_risk(rid: str, user: dict = Depends(get_current_user)):
    await db.risks.delete_one({"id": rid})
    return {"ok": True}

# ---------- Reviews ----------
@api.get("/reviews")
async def list_reviews(user: dict = Depends(get_current_user)):
    return await db.reviews.find({}, {"_id": 0}).sort("created_at", -1).to_list(500)

@api.post("/reviews")
async def create_review(body: ReviewIn, user: dict = Depends(get_current_user)):
    count = await db.reviews.count_documents({})
    r = Review(code=f"AR-{1000 + count + 1}", **body.model_dump())
    await db.reviews.insert_one(r.model_dump())
    return r.model_dump()

@api.put("/reviews/{rid}/stage")
async def update_review_stage(rid: str, stage: str, user: dict = Depends(get_current_user)):
    valid = ["Draft", "Submitted", "Domain Reviews", "EA Review", "Committee", "Approved", "Rejected"]
    if stage not in valid:
        raise HTTPException(status_code=400, detail="Invalid stage")
    await db.reviews.update_one({"id": rid}, {"$set": {"stage": stage}})
    return await db.reviews.find_one({"id": rid}, {"_id": 0})

@api.delete("/reviews/{rid}")
async def delete_review(rid: str, user: dict = Depends(get_current_user)):
    await db.reviews.delete_one({"id": rid})
    return {"ok": True}

# ---------- Dashboard KPIs ----------
@api.get("/kpis")
async def kpis(user: dict = Depends(get_current_user)):
    async def cnt(**q):
        return await db.objects.count_documents(q)
    return {
        "total_capabilities": await cnt(domain="business", type="capability"),
        "total_processes": await cnt(domain="business", type="process"),
        "critical_processes": await cnt(domain="business", type="process", criticality={"$in": ["High", "Mission-Critical"]}),
        "total_applications": await cnt(domain="application"),
        "mission_critical_apps": await cnt(domain="application", criticality="Mission-Critical"),
        "total_data_entities": await cnt(domain="data"),
        "total_integrations": await cnt(domain="integration"),
        "critical_integrations": await cnt(domain="integration", criticality={"$in": ["High", "Mission-Critical"]}),
        "total_security_controls": await cnt(domain="security"),
        "total_technologies": await cnt(domain="technology"),
        "eol_technologies": await cnt(domain="technology", status="Deprecated"),
        "open_risks": await db.risks.count_documents({"status": "Open"}),
        "pending_reviews": await db.reviews.count_documents({"stage": {"$in": ["Submitted", "Domain Reviews", "EA Review", "Committee"]}}),
        "total_adrs": await db.adrs.count_documents({}),
        "domain_counts": {d: await cnt(domain=d) for d in DOMAINS},
    }

# ---------- Audit ----------
async def audit(user: dict, action: str, obj_id: str, domain: str, before: Any, after: Any):
    await db.audit.insert_one({
        "id": str(uuid.uuid4()),
        "user_id": user.get("id"),
        "user_email": user.get("email"),
        "action": action,
        "object_id": obj_id,
        "domain": domain,
        "before": before,
        "after": after,
        "timestamp": datetime.now(timezone.utc).isoformat(),
    })

@api.get("/audit")
async def list_audit(
    domain: Optional[str] = None,
    user_email: Optional[str] = None,
    action: Optional[str] = None,
    q: Optional[str] = None,
    limit: int = 500,
    user: dict = Depends(get_current_user),
):
    query: Dict[str, Any] = {}
    if domain: query["domain"] = domain
    if user_email: query["user_email"] = user_email
    if action: query["action"] = action
    if q:
        query["$or"] = [
            {"user_email": {"$regex": q, "$options": "i"}},
            {"object_id": {"$regex": q, "$options": "i"}},
            {"action": {"$regex": q, "$options": "i"}},
        ]
    return await db.audit.find(query, {"_id": 0}).sort("timestamp", -1).limit(min(limit, 2000)).to_list(min(limit, 2000))

@api.get("/audit/users")
async def audit_users(user: dict = Depends(get_current_user)):
    users = await db.audit.distinct("user_email")
    return [u for u in users if u]

# ---------- Object Storage (Emergent or Local FS fallback) ----------
STORAGE_BASE = (os.environ.get("INTEGRATION_PROXY_URL") or "").strip() or "https://integrations.emergentagent.com"
STORAGE_URL = STORAGE_BASE.rstrip("/") + "/objstore/api/v1/storage"
EMERGENT_KEY = os.environ.get("EMERGENT_LLM_KEY")
APP_STORAGE_PREFIX = os.environ.get("APP_STORAGE_PREFIX", "colecle-eams")
STORAGE_MODE = os.environ.get("STORAGE_MODE", "auto").lower()  # auto | emergent | local
LOCAL_STORAGE_DIR = os.environ.get("LOCAL_STORAGE_DIR", "/data/uploads")
_storage_key = {"value": None}

def _use_local_storage() -> bool:
    if STORAGE_MODE == "local": return True
    if STORAGE_MODE == "emergent": return False
    return not EMERGENT_KEY  # auto: fall back to local when no Emergent key

def init_storage(force: bool = False):
    if _use_local_storage():
        os.makedirs(LOCAL_STORAGE_DIR, exist_ok=True)
        return "local"
    if _storage_key["value"] and not force:
        return _storage_key["value"]
    if not EMERGENT_KEY:
        return None
    r = requests.post(f"{STORAGE_URL}/init", json={"emergent_key": EMERGENT_KEY}, timeout=30)
    r.raise_for_status()
    _storage_key["value"] = r.json()["storage_key"]
    return _storage_key["value"]

def _storage_put(path: str, data: bytes, content_type: str) -> dict:
    if _use_local_storage():
        from onprem_storage import write_blob
        return write_blob(LOCAL_STORAGE_DIR, path, data)
    key = init_storage()
    if not key: raise HTTPException(status_code=503, detail="Object storage not configured")
    r = requests.put(f"{STORAGE_URL}/objects/{path}", headers={"X-Storage-Key": key, "Content-Type": content_type}, data=data, timeout=120)
    if r.status_code == 404:
        key = init_storage(force=True)
        r = requests.put(f"{STORAGE_URL}/objects/{path}", headers={"X-Storage-Key": key, "Content-Type": content_type}, data=data, timeout=120)
    r.raise_for_status()
    return r.json()

def _storage_get(path: str):
    if _use_local_storage():
        from onprem_storage import read_blob
        blob = read_blob(LOCAL_STORAGE_DIR, path)
        if blob is None:
            raise HTTPException(status_code=404, detail="File not found on disk")
        return blob, "application/octet-stream"
    key = init_storage()
    if not key: raise HTTPException(status_code=503, detail="Object storage not configured")
    r = requests.get(f"{STORAGE_URL}/objects/{path}", headers={"X-Storage-Key": key}, timeout=60)
    if r.status_code == 404:
        key = init_storage(force=True)
        r = requests.get(f"{STORAGE_URL}/objects/{path}", headers={"X-Storage-Key": key}, timeout=60)
    r.raise_for_status()
    return r.content, r.headers.get("Content-Type", "application/octet-stream")

# ---------- Documents (S3-style via Emergent Object Storage) ----------
@api.get("/objects/{oid}/documents")
async def list_documents(oid: str, user: dict = Depends(get_current_user)):
    return await db.documents.find({"object_id": oid, "is_deleted": {"$ne": True}}, {"_id": 0}).sort("uploaded_at", -1).to_list(200)

@api.post("/objects/{oid}/documents")
async def upload_document(oid: str, file: UploadFile = File(...), user: dict = Depends(get_current_user)):
    obj = await db.objects.find_one({"id": oid}, {"_id": 0})
    if not obj:
        raise HTTPException(status_code=404, detail="Object not found")
    data = await file.read()
    if len(data) > 200 * 1024 * 1024:
        raise HTTPException(status_code=413, detail="File too large (max 200MB)")
    ext = (file.filename.rsplit(".", 1)[-1] if "." in (file.filename or "") else "bin").lower()
    doc_id = str(uuid.uuid4())
    path = f"{APP_STORAGE_PREFIX}/objects/{oid}/{doc_id}.{ext}"
    ct = file.content_type or "application/octet-stream"
    try:
        result = _storage_put(path, data, ct)
    except HTTPException:
        raise
    except Exception as e:
        logger.error(f"Object storage upload failed: {e}")
        raise HTTPException(status_code=502, detail="Storage upload failed")
    doc = {
        "id": doc_id, "object_id": oid,
        "storage_path": result.get("path", path),
        "filename": file.filename, "content_type": ct,
        "size": result.get("size", len(data)),
        "uploaded_by": user.get("email"),
        "uploaded_at": datetime.now(timezone.utc).isoformat(),
        "is_deleted": False,
    }
    await db.documents.insert_one(doc)
    await audit(user, "upload_document", oid, obj.get("domain", ""), None, {"filename": doc["filename"], "size": doc["size"]})
    return {k: v for k, v in doc.items() if k != "_id"}

@api.get("/documents/{did}/download")
async def download_document(did: str, token: Optional[str] = None):
    if not token:
        raise HTTPException(status_code=401, detail="Not authenticated")
    try:
        payload = jwt.decode(token, JWT_SECRET, algorithms=[JWT_ALGO])
        u = await db.users.find_one({"id": payload["sub"]}, {"_id": 0})
        if not u: raise HTTPException(status_code=401, detail="Invalid")
    except HTTPException: raise
    except Exception:
        raise HTTPException(status_code=401, detail="Invalid token")
    doc = await db.documents.find_one({"id": did, "is_deleted": {"$ne": True}}, {"_id": 0})
    if not doc:
        raise HTTPException(status_code=404, detail="Not found")
    try:
        binary, ct = _storage_get(doc["storage_path"])
    except Exception as e:
        logger.error(f"Storage download failed: {e}")
        raise HTTPException(status_code=502, detail="Storage download failed")
    return StreamingResponse(
        io.BytesIO(binary),
        media_type=doc.get("content_type", ct),
        headers={"Content-Disposition": f'attachment; filename="{doc["filename"]}"'},
    )

@api.delete("/documents/{did}")
async def delete_document(did: str, user: dict = Depends(get_current_user)):
    doc = await db.documents.find_one({"id": did}, {"_id": 0})
    if not doc:
        raise HTTPException(status_code=404, detail="Not found")
    # Soft-delete (storage has no delete API)
    await db.documents.update_one({"id": did}, {"$set": {"is_deleted": True}})
    await audit(user, "delete_document", doc["object_id"], "", {"filename": doc["filename"]}, None)
    return {"ok": True}

# ---------- Process Flow (steps + per-step domain touches) ----------
@api.get("/processes/{pid}/flow")
async def process_flow(pid: str, user: dict = Depends(get_current_user)):
    process = await db.objects.find_one({"id": pid}, {"_id": 0})
    if not process:
        raise HTTPException(status_code=404, detail="Not found")
    step_rels = await db.relationships.find({"source_id": pid, "rel_type": "has_step"}, {"_id": 0}).to_list(200)
    step_ids = [r["target_id"] for r in step_rels]
    steps = await db.objects.find({"id": {"$in": step_ids}}, {"_id": 0}).to_list(200)
    # sort by code (PS-01, PS-02…)
    steps.sort(key=lambda s: (s.get("code") or "", s.get("name") or ""))
    # for each step, walk downstream 3 levels to gather touches per domain
    all_rels = await db.relationships.find({}, {"_id": 0}).to_list(20000)
    all_objs = await db.objects.find({}, {"_id": 0}).to_list(10000)
    by_id = {o["id"]: o for o in all_objs}
    out = []
    for s in steps:
        visited = {s["id"]}
        frontier = [(s["id"], 0)]
        while frontier:
            cur, d = frontier.pop(0)
            if d >= 3: continue
            for r in all_rels:
                if r["source_id"] == cur and r["target_id"] not in visited:
                    visited.add(r["target_id"])
                    frontier.append((r["target_id"], d + 1))
        touches: Dict[str, List[Dict[str, Any]]] = {}
        for nid in visited:
            if nid == s["id"]: continue
            n = by_id.get(nid)
            if not n: continue
            touches.setdefault(n["domain"], []).append(n)
        out.append({"step": s, "touches": touches})
    return {"process": process, "steps": out}

# ---------- Reports ----------
def _report_rows(kind: str, objs: List[dict], adrs: List[dict], stds: List[dict], risks: List[dict]):
    if kind == "application-portfolio":
        rows = [["Code", "Name", "Owner", "Criticality", "Lifecycle", "Status", "Technology Stack"]]
        for o in [x for x in objs if x["domain"] == "application"]:
            rows.append([o.get("code", ""), o.get("name", ""), o.get("owner", ""), o.get("criticality", ""), o.get("lifecycle", ""), o.get("status", ""), (o.get("attributes") or {}).get("tech", "")])
        return rows
    if kind == "technology-eol":
        rows = [["Code", "Name", "Vendor", "EOL Date", "Status", "Criticality"]]
        for o in [x for x in objs if x["domain"] == "technology"]:
            a = o.get("attributes") or {}
            rows.append([o.get("code", ""), o.get("name", ""), a.get("vendor", ""), a.get("eol", ""), o.get("status", ""), o.get("criticality", "")])
        return rows
    if kind == "business-capabilities":
        rows = [["Code", "Name", "Owner", "Criticality", "Maturity"]]
        for o in [x for x in objs if x["domain"] == "business" and x["type"] == "capability"]:
            a = o.get("attributes") or {}
            rows.append([o.get("code", ""), o.get("name", ""), o.get("owner", ""), o.get("criticality", ""), a.get("maturity", "")])
        return rows
    if kind == "integration-catalogue":
        rows = [["Code", "Name", "Owner", "Protocol", "Criticality", "Status"]]
        for o in [x for x in objs if x["domain"] == "integration"]:
            a = o.get("attributes") or {}
            rows.append([o.get("code", ""), o.get("name", ""), o.get("owner", ""), a.get("protocol", ""), o.get("criticality", ""), o.get("status", "")])
        return rows
    if kind == "data-catalogue":
        rows = [["Code", "Name", "Owner", "Classification", "Criticality"]]
        for o in [x for x in objs if x["domain"] == "data"]:
            a = o.get("attributes") or {}
            rows.append([o.get("code", ""), o.get("name", ""), o.get("owner", ""), a.get("classification", ""), o.get("criticality", "")])
        return rows
    if kind == "security-controls":
        rows = [["Code", "Name", "Owner", "Criticality", "Status"]]
        for o in [x for x in objs if x["domain"] == "security"]:
            rows.append([o.get("code", ""), o.get("name", ""), o.get("owner", ""), o.get("criticality", ""), o.get("status", "")])
        return rows
    if kind == "risks":
        rows = [["Code", "Title", "Domain", "Severity", "Owner", "Status"]]
        for r in risks:
            rows.append([r.get("code", ""), r.get("title", ""), r.get("domain", ""), r.get("severity", ""), r.get("owner", ""), r.get("status", "")])
        return rows
    if kind == "adrs":
        rows = [["Number", "Title", "Domain", "Status", "Owner", "Decision"]]
        for a in adrs:
            rows.append([f"ADR-{a.get('number'):03d}" if a.get("number") else "", a.get("title", ""), a.get("domain", ""), a.get("status", ""), a.get("owner", ""), a.get("decision", "")])
        return rows
    if kind == "standards":
        rows = [["Code", "Name", "Domain", "Mandatory", "Version", "Status"]]
        for s in stds:
            rows.append([s.get("code", ""), s.get("name", ""), s.get("domain", ""), "Yes" if s.get("mandatory") else "No", s.get("version", ""), s.get("status", "")])
        return rows
    if kind == "all-objects":
        rows = [["Code", "Name", "Domain", "Type", "Owner", "Criticality", "Status"]]
        for o in objs:
            rows.append([o.get("code", ""), o.get("name", ""), o.get("domain", ""), o.get("type", ""), o.get("owner", ""), o.get("criticality", ""), o.get("status", "")])
        return rows
    return [["No data"]]

REPORT_TITLES = {
    "application-portfolio": "Application Portfolio",
    "technology-eol": "Technology EOL / Lifecycle",
    "business-capabilities": "Business Capabilities",
    "integration-catalogue": "Integration Catalogue",
    "data-catalogue": "Data Catalogue",
    "security-controls": "Security Controls",
    "risks": "Architecture Risks",
    "adrs": "Architecture Decision Records",
    "standards": "Architecture Standards",
    "all-objects": "All Architecture Objects",
}

@api.get("/reports")
async def list_reports(user: dict = Depends(get_current_user)):
    return [{"key": k, "title": v} for k, v in REPORT_TITLES.items()]

@api.get("/reports/{kind}")
async def download_report(kind: str, format: str = "csv", user: dict = Depends(get_current_user)):
    if kind not in REPORT_TITLES:
        raise HTTPException(status_code=404, detail="Unknown report")
    objs = await db.objects.find({}, {"_id": 0}).to_list(10000)
    adrs = await db.adrs.find({}, {"_id": 0}).to_list(1000)
    stds = await db.standards.find({}, {"_id": 0}).to_list(1000)
    risks = await db.risks.find({}, {"_id": 0}).to_list(1000)
    rows = _report_rows(kind, objs, adrs, stds, risks)
    title = REPORT_TITLES[kind]
    fname_base = kind.replace("-", "_")

    if format == "csv":
        buf = io.StringIO()
        w = csv.writer(buf)
        w.writerows(rows)
        return StreamingResponse(
            io.BytesIO(buf.getvalue().encode()),
            media_type="text/csv",
            headers={"Content-Disposition": f'attachment; filename="{fname_base}.csv"'},
        )
    if format == "xlsx":
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        ws.title = title[:30]
        for row in rows:
            ws.append(row)
        # bold header
        from openpyxl.styles import Font, PatternFill
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F2937")
        for col in ws.columns:
            length = max(len(str(c.value or "")) for c in col)
            ws.column_dimensions[col[0].column_letter].width = min(length + 2, 60)
        out = io.BytesIO(); wb.save(out); out.seek(0)
        return StreamingResponse(
            out,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f'attachment; filename="{fname_base}.xlsx"'},
        )
    if format == "pdf":
        from reportlab.lib.pagesizes import A4, landscape
        from reportlab.lib import colors
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
        buf = io.BytesIO()
        doc = SimpleDocTemplate(buf, pagesize=landscape(A4), leftMargin=24, rightMargin=24, topMargin=32, bottomMargin=24)
        styles = getSampleStyleSheet()
        elements = [
            Paragraph(f"<b>{title}</b>", styles["Title"]),
            Paragraph(f"Colecle System EAMS · Generated {datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M UTC')}", styles["Normal"]),
            Spacer(1, 12),
        ]
        t = Table(rows, repeatRows=1)
        t.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F2937")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, -1), 8),
            ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#94a3b8")),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f8fafc")]),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ]))
        elements.append(t)
        doc.build(elements)
        buf.seek(0)
        return StreamingResponse(
            buf,
            media_type="application/pdf",
            headers={"Content-Disposition": f'attachment; filename="{fname_base}.pdf"'},
        )
    raise HTTPException(status_code=400, detail="Unknown format")

# ---------- Health ----------
@api.get("/")
async def root():
    return {"service": "Colecle EAMS", "status": "ok"}

# ---------- Seed ----------
async def seed_admin_and_data():
    await db.users.create_index("email", unique=True)
    await db.objects.create_index("id", unique=True)
    await db.objects.create_index([("domain", 1), ("type", 1)])
    await db.relationships.create_index("id", unique=True)

    admin_email = os.environ.get("ADMIN_EMAIL", "admin@example.com").lower()
    admin_pw = os.environ.get("ADMIN_PASSWORD", "admin123")
    existing = await db.users.find_one({"email": admin_email})
    if not existing:
        uid = str(uuid.uuid4())
        await db.users.insert_one({
            "id": uid, "email": admin_email, "password_hash": hash_pw(admin_pw),
            "name": "Chief Architect", "role": "admin",
            "created_at": datetime.now(timezone.utc).isoformat(),
        })
        logger.info(f"Seeded admin {admin_email}")
    elif not verify_pw(admin_pw, existing["password_hash"]):
        await db.users.update_one({"email": admin_email}, {"$set": {"password_hash": hash_pw(admin_pw)}})

    # Brand default
    if not await db.brand.find_one({"_id": "singleton"}):
        await db.brand.insert_one({"_id": "singleton", **BrandSettings().model_dump()})

    # Seed sample only if empty
    if await db.objects.count_documents({}) > 0:
        return

    def mk(domain, type, code, name, description="", **kw):
        obj = ArchObject(domain=domain, type=type, code=code, name=name, description=description, **kw).model_dump()
        return obj

    seed_objects = [
        # Business
        mk("business", "unit", "BU-01", "Payments Business Unit", "Owns payment products and services.", owner="COO", criticality="Mission-Critical"),
        mk("business", "capability", "CAP-01", "Payments", "Core capability enabling payment services.", owner="Head of Payments", criticality="Mission-Critical", attributes={"maturity": "L3", "strategic": "High"}),
        mk("business", "capability", "CAP-02", "Payment Processing", "Sub-capability under Payments.", owner="Head of Payments", criticality="Mission-Critical", attributes={"maturity": "L2", "strategic": "High"}),
        mk("business", "capability", "CAP-03", "Instant Payment Processing", "Real-time payment execution.", owner="Payments Ops", criticality="Mission-Critical", attributes={"maturity": "L2"}),
        mk("business", "process", "BP-01", "Instant Payment Execution", "End-to-end instant payment flow.", owner="Ops Lead", criticality="Mission-Critical", attributes={"rto": "5m", "rpo": "0s"}),
        mk("business", "process_step", "PS-01", "Receive Payment Request", "Accepts inbound payment instruction.", owner="Ops"),
        mk("business", "process_step", "PS-02", "Validate Participant", "Validates payer/payee.", owner="Ops"),
        mk("business", "process_step", "PS-03", "Process Transaction", "Executes debit/credit legs.", owner="Ops"),
        mk("business", "process_step", "PS-04", "Persist Transaction", "Persists to system of record.", owner="Ops"),
        mk("business", "process_step", "PS-05", "Send Response", "Responds to participant.", owner="Ops"),
        # Applications
        mk("application", "application", "APP-01", "Payment Gateway", "Public-facing gateway.", owner="App Team A", criticality="Mission-Critical", lifecycle="Production", attributes={"vendor": "Internal", "tech": "Java/Spring", "rto": "5m", "rpo": "0s", "dr_site": "DC-B", "has_dr": True}),
        mk("application", "application", "APP-02", "Core Payment Engine", "Executes payment logic.", owner="App Team A", criticality="Mission-Critical", lifecycle="Production", attributes={"tech": "Java/Spring", "rto": "1h", "rpo": "5m", "dr_site": "DC-B", "has_dr": True}),
        mk("application", "application", "APP-03", "Participant Master", "Master data for participants.", owner="MDM Team", criticality="High", lifecycle="Production", attributes={"rto": "4h", "rpo": "1h"}),
        # Data
        mk("data", "data_entity", "DE-01", "Payment Instruction", "Inbound instruction payload.", owner="Data Steward A", criticality="High", attributes={"classification": "Confidential"}),
        mk("data", "data_entity", "DE-02", "Participant Master Data", "Participant reference data.", owner="Data Steward B", criticality="High", attributes={"classification": "Internal"}),
        mk("data", "data_entity", "DE-03", "Payment Transaction", "System of record transaction.", owner="Data Steward A", criticality="Mission-Critical", attributes={"classification": "Restricted"}),
        mk("data", "data_store", "DS-01", "Payments Database", "Primary transactional DB.", owner="DBA", criticality="Mission-Critical", attributes={"engine": "PostgreSQL 15"}),
        # Integration
        mk("integration", "integration", "INT-01", "Public Payment REST API", "External inbound REST API.", owner="Int Team", criticality="Mission-Critical", attributes={"protocol": "REST/HTTPS"}),
        mk("integration", "integration", "INT-02", "Internal Participant API", "Internal service API.", owner="Int Team", criticality="High", attributes={"protocol": "REST"}),
        mk("integration", "integration", "INT-03", "API Gateway", "Central API gateway.", owner="Platform", criticality="Mission-Critical"),
        # Security
        mk("security", "control", "SEC-01", "TLS 1.3 In Transit", "Transport encryption.", owner="Security Team", criticality="Mission-Critical"),
        mk("security", "control", "SEC-02", "OAuth2 Authentication", "Authentication layer.", owner="Security Team", criticality="High"),
        mk("security", "control", "SEC-03", "RBAC Authorization", "Fine-grained authorization.", owner="Security Team", criticality="High"),
        # Technology
        mk("technology", "technology", "TECH-01", "API Gateway (Kong)", "Kong Gateway platform.", owner="Platform", criticality="Mission-Critical", attributes={"vendor": "Kong", "eol": "2027-06-01"}),
        mk("technology", "technology", "TECH-02", "Kubernetes Cluster", "OpenShift cluster.", owner="Platform", criticality="Mission-Critical", attributes={"vendor": "Red Hat"}),
        mk("technology", "technology", "TECH-03", "PostgreSQL 15", "Relational database engine.", owner="DBA", criticality="Mission-Critical", attributes={"vendor": "PostgreSQL"}),
        mk("technology", "infrastructure", "INF-01", "Primary Data Centre", "DC-A active site.", owner="Infra", criticality="Mission-Critical"),
    ]
    await db.objects.insert_many(seed_objects)

    # relationship helper by code
    by_code = {o["code"]: o["id"] for o in seed_objects}
    def rel(src, tgt, t="supports", crit="High"):
        return Relationship(source_id=by_code[src], target_id=by_code[tgt], rel_type=t, criticality=crit).model_dump()

    seed_rels = [
        # Business hierarchy
        rel("BU-01", "CAP-01", "owns"),
        rel("CAP-01", "CAP-02", "parent_of"),
        rel("CAP-02", "CAP-03", "parent_of"),
        rel("CAP-03", "BP-01", "realized_by", "Mission-Critical"),
        # Process steps
        rel("BP-01", "PS-01", "has_step"),
        rel("BP-01", "PS-02", "has_step"),
        rel("BP-01", "PS-03", "has_step"),
        rel("BP-01", "PS-04", "has_step"),
        rel("BP-01", "PS-05", "has_step"),
        # Steps -> Apps
        rel("PS-01", "APP-01", "uses_application"),
        rel("PS-02", "APP-03", "uses_application"),
        rel("PS-03", "APP-02", "uses_application"),
        rel("PS-04", "APP-02", "uses_application"),
        rel("PS-05", "APP-01", "uses_application"),
        # Apps -> Data
        rel("APP-01", "DE-01", "consumes_data"),
        rel("APP-03", "DE-02", "produces_data"),
        rel("APP-02", "DE-03", "produces_data", "Mission-Critical"),
        rel("DE-03", "DS-01", "stored_in", "Mission-Critical"),
        # Apps -> Integrations
        rel("APP-01", "INT-01", "exposes"),
        rel("APP-01", "INT-03", "routed_via"),
        rel("APP-03", "INT-02", "exposes"),
        # Integrations -> Security
        rel("INT-01", "SEC-01", "protected_by", "Mission-Critical"),
        rel("INT-01", "SEC-02", "protected_by"),
        rel("INT-02", "SEC-03", "protected_by"),
        # Apps -> Technology
        rel("APP-01", "TECH-01", "hosted_on"),
        rel("APP-02", "TECH-02", "hosted_on", "Mission-Critical"),
        rel("APP-03", "TECH-02", "hosted_on"),
        rel("DS-01", "TECH-03", "runs_on", "Mission-Critical"),
        rel("TECH-01", "INF-01", "hosted_in"),
        rel("TECH-02", "INF-01", "hosted_in"),
        rel("TECH-03", "INF-01", "hosted_in", "Mission-Critical"),
    ]
    await db.relationships.insert_many(seed_rels)

    # Standards
    await db.standards.insert_many([
        Standard(code="STD-001", name="TLS 1.3 for all external endpoints", domain="security", mandatory=True, owner="Security Team").model_dump(),
        Standard(code="STD-002", name="OpenAPI 3.1 for REST APIs", domain="integration", mandatory=True, owner="Integration Team").model_dump(),
        Standard(code="STD-003", name="PostgreSQL as preferred RDBMS", domain="data", mandatory=False, owner="Data Team").model_dump(),
        Standard(code="STD-004", name="Container-first deployment (Kubernetes)", domain="technology", mandatory=True, owner="Platform").model_dump(),
    ])

    # ADRs
    await db.adrs.insert_many([
        ADR(number=1, title="Adopt Kong as API Gateway", domain="integration", context="Need central API gateway for payment services.", decision="Adopt Kong Gateway on OpenShift.", rationale="Mature OSS + enterprise support.", status="Accepted", owner="Chief Architect").model_dump(),
        ADR(number=2, title="Standardize on PostgreSQL 15", domain="data", context="Consolidate database engines.", decision="Standardize on PostgreSQL 15.", rationale="Feature parity + open source.", status="Accepted", owner="Data Architect").model_dump(),
    ])

    # Risks
    await db.risks.insert_many([
        Risk(code="RISK-001", title="Kong Gateway single point of failure", domain="integration", description="No multi-region failover configured.", likelihood="Medium", impact="High", score=12, severity="High", status="Open", owner="Platform").model_dump(),
        Risk(code="RISK-002", title="Legacy TLS versions on internal APIs", domain="security", description="Some internal APIs still support TLS 1.2.", likelihood="Low", impact="Medium", score=4, severity="Low", status="Open", owner="Security Team").model_dump(),
    ])

    # Reviews
    await db.reviews.insert_many([
        Review(code="AR-1001", title="Instant Payment Engine v2 - Architecture Review", submitter="Lead Architect", summary="Introduce v2 engine with real-time features.", domains_impacted=["application", "data", "integration", "security"], stage="EA Review").model_dump(),
    ])

    logger.info("Seeded EAMS sample data")

    # Write test creds
    try:
        mem_dir = Path("/app/memory")
        mem_dir.mkdir(exist_ok=True)
        (mem_dir / "test_credentials.md").write_text(
            f"# Test Credentials\n\n## Admin\n- Email: {admin_email}\n- Password: {admin_pw}\n- Role: admin\n\n## Endpoints\n- POST /api/auth/login\n- GET /api/auth/me\n- POST /api/auth/logout\n"
        )
    except Exception as e:
        logger.warning(f"Failed to write test_credentials.md: {e}")

@app.on_event("startup")
async def startup():
    await seed_admin_and_data()
    try:
        init_storage()
        mode = "local FS" if _use_local_storage() else "Emergent"
        logger.info(f"Object storage initialized ({mode})")
    except Exception as e:
        logger.warning(f"Object storage init failed (documents feature will be unavailable): {e}")
    # Optional in-process scheduler for on-prem deployments without external cron.
    if os.environ.get("RUN_CRONS", "false").lower() in ("1", "true", "yes"):
        try:
            from apscheduler.schedulers.asyncio import AsyncIOScheduler
            from apscheduler.triggers.cron import CronTrigger
            sched = AsyncIOScheduler(timezone="UTC")
            sched.add_job(
                lambda: _run_weekly_reports_job(str(uuid.uuid4())),
                CronTrigger(day_of_week="mon", hour=8, minute=0),
                id="weekly-reports",
                replace_existing=True,
            )
            sched.start()
            logger.info("APScheduler started: weekly-reports Mon 08:00 UTC")
        except Exception as e:
            logger.warning(f"APScheduler start failed: {e}")

# ---------- DR Coverage ----------
@api.get("/dr/coverage")
async def dr_coverage(user: dict = Depends(get_current_user)):
    apps = await db.objects.find({"domain": "application"}, {"_id": 0}).to_list(2000)
    covered, missing, partial = [], [], []
    for a in apps:
        attr = a.get("attributes") or {}
        has_dr = bool(attr.get("dr_site")) or attr.get("has_dr") is True
        rto = attr.get("rto")
        rpo = attr.get("rpo")
        entry = {**a, "_dr": {"dr_site": attr.get("dr_site"), "rto": rto, "rpo": rpo, "has_dr": has_dr}}
        if has_dr and rto and rpo:
            covered.append(entry)
        elif has_dr or rto or rpo:
            partial.append(entry)
        else:
            missing.append(entry)
    total = len(apps)
    return {
        "total": total,
        "covered": covered,
        "partial": partial,
        "missing": missing,
        "coverage_pct": round(100 * len(covered) / total, 1) if total else 0,
        "critical_missing": [x for x in missing if x.get("criticality") in ("Mission-Critical", "High")],
    }

# ---------- Capability Heatmap ----------
_MATURITY_LEVELS = ["L1", "L2", "L3", "L4", "L5"]
_STRATEGIC_LEVELS = ["Low", "Medium", "High"]

@api.get("/capabilities/heatmap")
async def capability_heatmap(user: dict = Depends(get_current_user)):
    caps = await db.objects.find({"domain": "business", "type": "capability"}, {"_id": 0}).to_list(2000)
    grid: Dict[str, Dict[str, List[Dict[str, Any]]]] = {m: {s: [] for s in _STRATEGIC_LEVELS} for m in _MATURITY_LEVELS}
    for c in caps:
        a = c.get("attributes") or {}
        m = (a.get("maturity") or "L1").upper()
        s = a.get("strategic") or "Medium"
        if m not in _MATURITY_LEVELS: m = "L1"
        if s not in _STRATEGIC_LEVELS: s = "Medium"
        grid[m][s].append(c)
    return {"grid": grid, "maturity_levels": _MATURITY_LEVELS, "strategic_levels": _STRATEGIC_LEVELS, "total": len(caps)}

# ---------- Report Subscriptions ----------
class SubscriptionIn(BaseModel):
    email: EmailStr
    report_keys: List[str]

@api.get("/subscriptions")
async def list_subscriptions(user: dict = Depends(get_current_user)):
    return await db.subscriptions.find({}, {"_id": 0}).sort("email", 1).to_list(500)

@api.post("/subscriptions")
async def create_subscription(body: SubscriptionIn, user: dict = Depends(get_current_user)):
    invalid = [k for k in body.report_keys if k not in REPORT_TITLES]
    if invalid: raise HTTPException(status_code=400, detail=f"Unknown report keys: {invalid}")
    sub = {
        "id": str(uuid.uuid4()),
        "email": body.email.lower(),
        "report_keys": body.report_keys,
        "created_by": user.get("email"),
        "created_at": datetime.now(timezone.utc).isoformat(),
        "last_sent_at": None,
    }
    await db.subscriptions.update_one({"email": sub["email"]}, {"$set": sub}, upsert=True)
    return sub

@api.delete("/subscriptions/{sid}")
async def delete_subscription(sid: str, user: dict = Depends(get_current_user)):
    await db.subscriptions.delete_one({"id": sid})
    return {"ok": True}

# ---------- Email helper (Emergent-managed Resend) ----------
EMAIL_BASE_URL = "https://integrations.emergentagent.com"
EMAIL_KEY = os.environ.get("EMERGENT_EMAIL_KEY")
EMAIL_FROM_NAME = os.environ.get("EMAIL_FROM_NAME", "Colecle EAMS")

async def send_report_email(to: str, subject: str, html: str) -> Optional[str]:
    if not EMAIL_KEY:
        logger.info(f"[EMAIL DRY-RUN] to={to} subject={subject!r}")
        return "dry-run"
    payload = {"to": [to], "subject": subject, "html": html, "from_name": EMAIL_FROM_NAME}
    try:
        async with httpx.AsyncClient(timeout=30) as client:
            r = await client.post(f"{EMAIL_BASE_URL}/api/v1/email/send", headers={"X-Email-Key": EMAIL_KEY}, json=payload)
        r.raise_for_status()
        return r.json().get("id")
    except Exception as e:
        logger.error(f"Email send failed: {e}")
        return None

# ---------- Cron endpoint: Weekly reports ----------
WEBHOOK_CRON_SECRET = os.environ.get("WEBHOOK_CRON_SECRET", "")

async def _run_weekly_reports_job(run_id: str):
    # Idempotency: check by run_id
    if run_id and await db.cron_runs.find_one({"run_id": run_id}):
        return
    if run_id:
        await db.cron_runs.insert_one({"run_id": run_id, "job": "weekly-reports", "at": datetime.now(timezone.utc).isoformat()})
    subs = await db.subscriptions.find({}, {"_id": 0}).to_list(2000)
    for s in subs:
        titles = [REPORT_TITLES[k] for k in s.get("report_keys", []) if k in REPORT_TITLES]
        if not titles: continue
        subject = f"Your weekly Colecle EAMS reports · {datetime.now(timezone.utc).strftime('%b %d, %Y')}"
        rows_html = "".join(f"<li style='margin:6px 0'>{escape_html(t)}</li>" for t in titles)
        html = (
            "<table role='presentation' width='100%' style='font-family:Arial,sans-serif;color:#111'>"
            "<tr><td style='padding:24px'>"
            f"<h2 style='color:#f59e0b;margin:0 0 8px'>Colecle EAMS Weekly Digest</h2>"
            f"<p>Hello,</p><p>Here are your subscribed architecture reports for this week:</p>"
            f"<ul>{rows_html}</ul>"
            f"<p>Sign in to Colecle EAMS to download each report as PDF, Excel or CSV.</p>"
            f"<p style='font-size:12px;color:#888;margin-top:24px'>Sent by {escape_html(EMAIL_FROM_NAME)}. "
            "We never ask for your password by email. Manage your subscription in Reports & Exports → Subscriptions.</p>"
            "</td></tr></table>"
        )
        eid = await send_report_email(s["email"], subject, html)
        await db.subscriptions.update_one({"id": s["id"]}, {"$set": {"last_sent_at": datetime.now(timezone.utc).isoformat(), "last_email_id": eid}})

def escape_html(s: str) -> str:
    return (s or "").replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")

@api.post("/cron/weekly-reports")
async def cron_weekly_reports(request: Request):
    # Cron endpoints must ack 2xx immediately; enqueue/background the actual work.
    auth = request.headers.get("Authorization", "")
    if not auth.startswith("Bearer "):
        raise HTTPException(status_code=401, detail="Missing auth")
    token = auth[7:]
    if not WEBHOOK_CRON_SECRET or not hmac.compare_digest(token, WEBHOOK_CRON_SECRET):
        raise HTTPException(status_code=401, detail="Invalid cron secret")
    body = {}
    try:
        body = await request.json()
    except Exception:
        pass
    run_id = request.headers.get("X-Webhook-Id") or body.get("run_id") or str(uuid.uuid4())
    import asyncio
    asyncio.create_task(_run_weekly_reports_job(run_id))
    return {"ok": True, "queued": True, "run_id": run_id}

@api.post("/subscriptions/{sid}/send-now")
async def send_now(sid: str, user: dict = Depends(get_current_user)):
    sub = await db.subscriptions.find_one({"id": sid}, {"_id": 0})
    if not sub:
        raise HTTPException(status_code=404, detail="Not found")
    titles = [REPORT_TITLES[k] for k in sub.get("report_keys", []) if k in REPORT_TITLES]
    subject = f"Colecle EAMS reports · {datetime.now(timezone.utc).strftime('%b %d, %Y')}"
    rows_html = "".join(f"<li>{escape_html(t)}</li>" for t in titles)
    html = (
        "<table role='presentation' width='100%' style='font-family:Arial,sans-serif;color:#111'>"
        f"<tr><td style='padding:24px'><h2 style='color:#f59e0b'>Colecle EAMS Digest</h2>"
        f"<ul>{rows_html}</ul>"
        f"<p style='font-size:12px;color:#888'>Sent by {escape_html(EMAIL_FROM_NAME)}. Manage subscription in Reports.</p></td></tr></table>"
    )
    eid = await send_report_email(sub["email"], subject, html)
    await db.subscriptions.update_one({"id": sid}, {"$set": {"last_sent_at": datetime.now(timezone.utc).isoformat(), "last_email_id": eid}})
    return {"ok": True, "email_id": eid, "dry_run": eid == "dry-run"}

# ---------- LDAP / Directory Sign-in (admin-configurable) ----------
LDAP_ROLES = ["admin", "lead_architect", "domain_architect", "reviewer", "viewer"]

class LdapSettings(BaseModel):
    enabled: bool = False
    server_url: str = ""  # e.g. ldaps://ad.example.local:636
    start_tls: bool = False
    verify_cert: bool = True
    ca_cert_pem: Optional[str] = None  # optional CA chain
    bind_dn: str = ""  # service account
    bind_password: str = ""  # write-only
    user_search_base: str = ""
    user_filter: str = "(&(objectCategory=person)(objectClass=user)({attr}={username}))"
    login_attribute: str = "sAMAccountName"
    email_attribute: str = "mail"
    name_attribute: str = "displayName"
    default_role: str = "viewer"
    group_role_mappings: List[Dict[str, str]] = []  # [{"group_dn": "CN=EAMS-Admins,...", "role": "admin"}]
    connect_timeout: int = 5
    updated_at: str = Field(default_factory=lambda: datetime.now(timezone.utc).isoformat())

async def get_ldap_settings() -> dict:
    doc = await db.ldap_settings.find_one({"_id": "singleton"}, {"_id": 0})
    if not doc:
        doc = LdapSettings().model_dump()
        await db.ldap_settings.insert_one({"_id": "singleton", **doc})
    return doc

def _mask_ldap(cfg: dict) -> dict:
    out = {k: v for k, v in cfg.items() if k != "bind_password"}
    out["bind_password_set"] = bool(cfg.get("bind_password"))
    return out

@api.get("/admin/ldap")
async def admin_get_ldap(user: dict = Depends(get_current_user)):
    if user.get("role") != "admin":
        raise HTTPException(status_code=403, detail="Admin only")
    cfg = await get_ldap_settings()
    return _mask_ldap(cfg)

@api.put("/admin/ldap")
async def admin_put_ldap(body: LdapSettings, user: dict = Depends(get_current_user)):
    if user.get("role") != "admin":
        raise HTTPException(status_code=403, detail="Admin only")
    cur = await get_ldap_settings()
    payload = body.model_dump()
    # Preserve stored password when caller sends empty
    if not payload.get("bind_password"):
        payload["bind_password"] = cur.get("bind_password", "")
    for m in payload.get("group_role_mappings", []):
        if m.get("role") not in LDAP_ROLES:
            raise HTTPException(status_code=400, detail=f"Invalid role: {m.get('role')}")
    payload["updated_at"] = datetime.now(timezone.utc).isoformat()
    await db.ldap_settings.update_one({"_id": "singleton"}, {"$set": payload}, upsert=True)
    await audit(user, "update_ldap_settings", "ldap", "admin", _mask_ldap(cur), _mask_ldap(payload))
    return _mask_ldap(payload)

def _build_ldap_server(cfg: dict):
    from ldap3 import Server, Tls
    import ssl, tempfile
    if not cfg.get("server_url"):
        raise HTTPException(status_code=400, detail="LDAP not configured")
    url = cfg["server_url"].strip()
    use_ssl = url.lower().startswith("ldaps://")
    tls = None
    if use_ssl or cfg.get("start_tls"):
        ca_file = None
        if cfg.get("ca_cert_pem"):
            f = tempfile.NamedTemporaryFile(mode="w", suffix=".pem", delete=False)
            f.write(cfg["ca_cert_pem"]); f.close()
            ca_file = f.name
        tls = Tls(
            validate=ssl.CERT_REQUIRED if cfg.get("verify_cert", True) else ssl.CERT_NONE,
            ca_certs_file=ca_file,
        )
    return Server(url, use_ssl=use_ssl, tls=tls, connect_timeout=int(cfg.get("connect_timeout") or 5), get_info=None)

def _ldap_authenticate_sync(cfg: dict, username: str, password: str) -> dict:
    from ldap3 import Connection, SIMPLE, SUBTREE
    from ldap3.utils.conv import escape_filter_chars
    from ldap3.core.exceptions import LDAPCommunicationError, LDAPSocketOpenError, LDAPBindError
    server = _build_ldap_server(cfg)
    safe_username = escape_filter_chars(username)
    attr = cfg.get("login_attribute", "sAMAccountName")
    user_filter = (cfg.get("user_filter") or "(&(objectCategory=person)(objectClass=user)({attr}={username}))").format(attr=attr, username=safe_username)
    email_attr = cfg.get("email_attribute", "mail")
    name_attr = cfg.get("name_attribute", "displayName")
    svc = None
    try:
        svc = Connection(
            server,
            user=cfg.get("bind_dn") or None,
            password=cfg.get("bind_password") or None,
            authentication=SIMPLE if cfg.get("bind_dn") else None,
            auto_bind=False, read_only=True, receive_timeout=int(cfg.get("connect_timeout") or 5),
        )
        if cfg.get("start_tls"): svc.start_tls()
        if not svc.bind():
            raise HTTPException(status_code=502, detail=f"LDAP service bind failed: {svc.last_error}")
        svc.search(
            cfg.get("user_search_base") or "",
            user_filter, search_scope=SUBTREE,
            attributes=["distinguishedName", attr, email_attr, name_attr, "memberOf"],
        )
        if len(svc.entries) != 1:
            raise HTTPException(status_code=401, detail="Invalid corporate credentials")
        e = svc.entries[0]
        user_dn = e.entry_dn
        # Bind AS THE USER to verify password (never trust svc account for auth check)
        user_conn = Connection(server, user=user_dn, password=password, authentication=SIMPLE, read_only=True, receive_timeout=int(cfg.get("connect_timeout") or 5))
        if cfg.get("start_tls"): user_conn.start_tls()
        if not user_conn.bind():
            raise HTTPException(status_code=401, detail="Invalid corporate credentials")
        try: groups = {str(g) for g in (e.memberOf.values if hasattr(e, "memberOf") else [])}
        except Exception: groups = set()
        try: email = str(getattr(e, email_attr).value) if hasattr(e, email_attr) else None
        except Exception: email = None
        try: display = str(getattr(e, name_attr).value) if hasattr(e, name_attr) else username
        except Exception: display = username
        # Role from group mappings (case-insensitive DN compare)
        role = cfg.get("default_role") or "viewer"
        role_priority = {"admin": 4, "lead_architect": 3, "domain_architect": 2, "reviewer": 1, "viewer": 0}
        gset_lower = {g.lower() for g in groups}
        for m in cfg.get("group_role_mappings", []):
            gdn = (m.get("group_dn") or "").lower()
            if gdn and gdn in gset_lower and role_priority.get(m.get("role"), 0) > role_priority.get(role, 0):
                role = m["role"]
        return {"user_dn": user_dn, "username": username.lower(), "email": email, "name": display, "groups": list(groups), "role": role}
    except HTTPException: raise
    except (LDAPCommunicationError, LDAPSocketOpenError, TimeoutError, OSError) as ex:
        raise HTTPException(status_code=503, detail=f"LDAP directory unreachable: {ex}")
    except LDAPBindError as ex:
        raise HTTPException(status_code=401, detail="Invalid corporate credentials")
    finally:
        try:
            if svc: svc.unbind()
        except Exception: pass

@api.post("/admin/ldap/test")
async def admin_ldap_test(body: Dict[str, Any], user: dict = Depends(get_current_user)):
    if user.get("role") != "admin":
        raise HTTPException(status_code=403, detail="Admin only")
    cfg = await get_ldap_settings()
    # Merge overrides but keep stored password if none provided
    for k, v in (body or {}).items():
        if v not in (None, ""):
            cfg[k] = v
    test_username = (body or {}).get("test_username")
    test_password = (body or {}).get("test_password")
    if not test_username:
        # Just verify service bind
        try:
            from ldap3 import Connection, SIMPLE
            server = _build_ldap_server(cfg)
            conn = Connection(server, user=cfg.get("bind_dn") or None, password=cfg.get("bind_password") or None,
                              authentication=SIMPLE if cfg.get("bind_dn") else None, read_only=True,
                              receive_timeout=int(cfg.get("connect_timeout") or 5))
            if cfg.get("start_tls"): conn.start_tls()
            ok = conn.bind()
            err = None if ok else str(conn.last_error)
            try: conn.unbind()
            except Exception: pass
            return {"ok": ok, "stage": "service_bind", "error": err}
        except HTTPException as e: return {"ok": False, "stage": "service_bind", "error": e.detail}
        except Exception as e: return {"ok": False, "stage": "service_bind", "error": str(e)}
    # Full end-to-end test
    try:
        import asyncio as _asyncio
        result = await _asyncio.to_thread(_ldap_authenticate_sync, cfg, test_username, test_password)
        return {"ok": True, "stage": "full", "resolved_user": {k: result[k] for k in ("username", "email", "name", "role", "user_dn")}, "groups_found": len(result.get("groups", []))}
    except HTTPException as e:
        return {"ok": False, "stage": "full", "error": e.detail}
    except Exception as e:
        return {"ok": False, "stage": "full", "error": str(e)}

# ---------- LDAP / Directory sign-in (public endpoint) ----------
LDAP_DOMAIN = "colecle.corp"
LDAP_USER_MAP = {
    # samAccountName / uid : local user email
}

class LdapLoginIn(BaseModel):
    username: str  # accepts "COLECLE\\user", "user@colecle.corp", or "user"
    password: str

@api.post("/auth/ldap")
async def ldap_login(body: LdapLoginIn, response: Response):
    raw = body.username.strip()
    if "\\" in raw:
        _, uname = raw.split("\\", 1)
    elif "@" in raw:
        uname = raw.split("@", 1)[0]
    else:
        uname = raw
    cfg = await get_ldap_settings()
    if cfg.get("enabled") and cfg.get("server_url"):
        # Real LDAP bind
        import asyncio as _asyncio
        result = await _asyncio.to_thread(_ldap_authenticate_sync, cfg, uname, body.password)
        # Upsert local user (LDAP-sourced), tracked by email or ldap:<username>
        email = (result.get("email") or f"{uname}@ldap.local").lower()
        existing = await db.users.find_one({"email": email})
        if existing:
            await db.users.update_one({"email": email}, {"$set": {"name": result["name"], "role": result["role"], "auth_source": "ldap", "last_login": datetime.now(timezone.utc).isoformat(), "ldap_dn": result["user_dn"]}})
            user = await db.users.find_one({"email": email}, {"_id": 0, "password_hash": 0})
        else:
            uid = str(uuid.uuid4())
            user = {
                "id": uid, "email": email, "name": result["name"], "role": result["role"],
                "password_hash": hash_pw(secrets.token_hex(24)),  # unusable random password
                "auth_source": "ldap", "ldap_dn": result["user_dn"],
                "created_at": datetime.now(timezone.utc).isoformat(),
                "last_login": datetime.now(timezone.utc).isoformat(),
            }
            await db.users.insert_one(user)
        token = make_token(user["id"], user["email"], user["role"])
        response.set_cookie("access_token", token, httponly=True, secure=True, samesite="none", max_age=60*60*24*7, path="/")
        return {"id": user["id"], "email": user["email"], "name": user["name"], "role": user["role"], "token": token, "auth_method": "ldap"}
    # Preview / not configured: local shadow fallback
    uname_l = uname.lower()
    user = await db.users.find_one({"email": {"$regex": f"^{uname_l}@", "$options": "i"}})
    if not user or not verify_pw(body.password, user["password_hash"]):
        raise HTTPException(status_code=401, detail="Invalid corporate credentials")
    token = make_token(user["id"], user["email"], user["role"])
    response.set_cookie("access_token", token, httponly=True, secure=True, samesite="none", max_age=60*60*24*7, path="/")
    return {"id": user["id"], "email": user["email"], "name": user["name"], "role": user["role"], "token": token, "auth_method": "ldap-preview"}

app.include_router(api)
app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=False,
    allow_methods=["*"],
    allow_headers=["*"],
)

@app.on_event("shutdown")
async def shutdown():
    client.close()
